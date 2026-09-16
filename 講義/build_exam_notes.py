# -*- coding: utf-8 -*-
"""敷地計畫與都市設計｜考試筆記（總整理）Word 版。
單一事實來源：敷地計畫與都市設計/23-考試筆記（總整理）.md
改 markdown 重跑本檔即更新。
"""
import re
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = "/home/user/architecture-study-notes/敷地計畫與都市設計/23-考試筆記（總整理）.md"
OUT = "/home/user/architecture-study-notes/講義/敷地計畫考試筆記_總整理.docx"

CJK = 'Microsoft JhengHei'
NAVY = RGBColor(0x1F, 0x3B, 0x5C)
RED = RGBColor(0xC0, 0x39, 0x2B)
GRAY = RGBColor(0x55, 0x55, 0x55)
GREEN = RGBColor(0x2E, 0x6B, 0x33)
BODY = 9.5

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(1.4)
sec.top_margin = Cm(1.2); sec.bottom_margin = Cm(1.2)

st = doc.styles['Normal']; st.font.name = CJK; st.font.size = Pt(BODY)
_rf = st.element.get_or_add_rPr().get_or_add_rFonts()
for a in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
    _rf.set(qn(a), CJK)


def _f(r):
    r.font.name = CJK
    rr = r._element.get_or_add_rPr().get_or_add_rFonts()
    for a in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
        rr.set(qn(a), CJK)


def shade(cell, hx):
    tc = cell._tc.get_or_add_tcPr()
    e = OxmlElement('w:shd'); e.set(qn('w:val'), 'clear'); e.set(qn('w:fill'), hx)
    tc.append(e)


def tight(p, before=0, after=1.2, line=1.06):
    pf = p.paragraph_format
    pf.space_before = Pt(before); pf.space_after = Pt(after); pf.line_spacing = line


INLINE = re.compile(r'(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))')


def emit(p, text, size=BODY, color=None, base_bold=False):
    """處理 **粗體**、`code`、[連結](url) 的行內標記"""
    for part in INLINE.split(text):
        if not part:
            continue
        bold, col, t = base_bold, color, part
        if part.startswith('**') and part.endswith('**'):
            bold, t = True, part[2:-2]
        elif part.startswith('`') and part.endswith('`'):
            t = part[1:-1]; col = col or NAVY
        else:
            m = re.fullmatch(r'\[([^\]]+)\]\([^)]+\)', part)
            if m:
                t = m.group(1); col = col or NAVY
        # 語意著色
        if col is None:
            if t.startswith('⚠') or t.startswith('🔴') or t.startswith('❌'):
                col = RED
            elif t.startswith('✅') or t.startswith('💡'):
                col = GREEN
        r = p.add_run(t); r.font.size = Pt(size); r.font.bold = bold
        if col:
            r.font.color.rgb = col
        _f(r)


def para(text, size=BODY, indent=0.0, before=0, after=1.2, color=None, bold=False):
    p = doc.add_paragraph(); tight(p, before, after)
    if indent:
        p.paragraph_format.left_indent = Cm(indent)
    emit(p, text, size, color, bold)
    return p


def heading(text, level):
    sizes = {1: 15, 2: 12.5, 3: 10.5}
    fills = {1: '1F3B5C', 2: '35617F'}
    text = re.sub(r'\s*\{#.*\}$', '', text)
    if level <= 2:
        t = doc.add_table(rows=1, cols=1); t.style = 'Table Grid'
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        c = t.rows[0].cells[0]; shade(c, fills[level]); c.text = ''
        p = c.paragraphs[0]; tight(p, 1, 1, 1.0)
        r = p.add_run(text.replace('★', '★ ')); r.font.size = Pt(sizes[level])
        r.font.bold = True; r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF); _f(r)
        doc.add_paragraph().paragraph_format.space_after = Pt(2)
    else:
        p = doc.add_paragraph(); tight(p, 5, 2)
        r = p.add_run('▌' + text); r.font.size = Pt(sizes[3])
        r.font.bold = True; r.font.color.rgb = NAVY; _f(r)


def table(rows):
    ncol = max(len(r) for r in rows)
    t = doc.add_table(rows=len(rows), cols=ncol); t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = True
    for i, row in enumerate(rows):
        for j in range(ncol):
            cell = t.rows[i].cells[j]
            cell.text = ''
            p = cell.paragraphs[0]; tight(p, 0.5, 0.5, 1.0)
            txt = row[j] if j < len(row) else ''
            if i == 0:
                shade(cell, 'E4EAF0')
                emit(p, txt, BODY - 0.5, NAVY, base_bold=True)
            else:
                emit(p, txt, BODY - 0.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def quote(lines):
    t = doc.add_table(rows=1, cols=1); t.style = 'Table Grid'
    c = t.rows[0].cells[0]; shade(c, 'FFF7E0'); c.text = ''
    first = True
    for ln in lines:
        p = c.paragraphs[0] if first else c.add_paragraph()
        tight(p, 0.5, 0.5, 1.04); first = False
        emit(p, ln, BODY - 0.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def code(lines):
    t = doc.add_table(rows=1, cols=1); t.style = 'Table Grid'
    c = t.rows[0].cells[0]; shade(c, 'F2F4F6'); c.text = ''
    first = True
    for ln in lines:
        p = c.paragraphs[0] if first else c.add_paragraph()
        tight(p, 0, 0, 1.0); first = False
        r = p.add_run(ln.replace('\t', '    ')); r.font.size = Pt(8.0)
        r.font.name = 'Consolas'
        rr = r._element.get_or_add_rPr().get_or_add_rFonts()
        rr.set(qn('w:ascii'), 'Consolas'); rr.set(qn('w:hAnsi'), 'Consolas')
        rr.set(qn('w:eastAsia'), CJK)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


# ---------------- 解析 ----------------
lines = open(SRC, encoding='utf-8').read().split('\n')
i = 0
buf_q, buf_c, buf_t = [], [], []
in_code = False
skip_toc = False


def flush():
    global buf_q, buf_t
    if buf_q:
        quote(buf_q); buf_q = []
    if buf_t:
        table(buf_t); buf_t = []


while i < len(lines):
    ln = lines[i].rstrip()
    i += 1

    if ln.startswith('```'):
        if in_code:
            code(buf_c); buf_c = []; in_code = False
        else:
            flush(); in_code = True
        continue
    if in_code:
        buf_c.append(ln); continue

    if ln.startswith('#'):
        flush()
        lvl = len(ln) - len(ln.lstrip('#'))
        heading(ln.lstrip('# ').strip(), min(lvl, 3))
        continue

    if ln.startswith('>'):
        if buf_t:
            table(buf_t); buf_t = []
        body = ln[1:].strip()
        if body:
            buf_q.append(body)
        continue
    if buf_q and not ln.startswith('>'):
        quote(buf_q); buf_q = []

    if ln.startswith('|'):
        cells = [c.strip() for c in ln.strip('|').split('|')]
        if all(re.fullmatch(r':?-{2,}:?', c) for c in cells if c):
            continue
        buf_t.append(cells)
        continue
    if buf_t:
        table(buf_t); buf_t = []

    if re.fullmatch(r'-{3,}', ln):
        p = doc.add_paragraph(); tight(p, 1, 3)
        pPr = p._p.get_or_add_pPr(); pbdr = OxmlElement('w:pBdr')
        bot = OxmlElement('w:bottom')
        bot.set(qn('w:val'), 'single'); bot.set(qn('w:sz'), '6')
        bot.set(qn('w:space'), '1'); bot.set(qn('w:color'), 'BBBBBB')
        pbdr.append(bot); pPr.append(pbdr)
        continue

    if not ln.strip():
        continue

    m = re.match(r'^(\s*)- \[ \] (.*)', ln)
    if m:
        para('☐　' + m.group(2), indent=0.5 + len(m.group(1)) * 0.25)
        continue
    m = re.match(r'^(\s*)[-*] (.*)', ln)
    if m:
        para('・' + m.group(2), indent=0.5 + len(m.group(1)) * 0.25)
        continue
    m = re.match(r'^(\s*)(\d+)\. (.*)', ln)
    if m:
        para(f'{m.group(2)}. {m.group(3)}', indent=0.5 + len(m.group(1)) * 0.25)
        continue

    para(ln.strip())

flush()
if buf_c:
    code(buf_c)

doc.save(OUT)
print("saved", OUT)
